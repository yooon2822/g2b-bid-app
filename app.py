import io
import math
import requests
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- scipy 의존성을 완전히 제거한 순수 수학 정규분포 함수 ---
def norm_cdf(x, mu=0.0, sigma=1.0):
    """표준 오차함수를 이용한 정규분포 누적확률(CDF) 계산"""
    return 0.5 * (1.0 + math.erf((x - mu) / (sigma * math.sqrt(2.0))))

def norm_ppf(p, mu=0.0, sigma=1.0):
    """정규분포 역함수(Quantile/PPF) 고정밀 근사 알고리즘 (Beasley-Springer-Moro)"""
    if p <= 0.0:
        return mu - 5.0 * sigma
    if p >= 1.0:
        return mu + 5.0 * sigma

    a = [2.50662823884, -18.61500062529, 41.39119773534, -25.44106049637]
    b = [-8.47351093090, 23.08336743743, -21.06224101826, 3.13082909833]
    c = [0.3374754822726147, 0.9761690190917186, 0.1607979714918209,
         0.02764388103300558, 0.0038405729373609, 0.0003951896511919,
         0.0000321767881768, 0.0000002888167364, 0.0000003960315187]

    y = p - 0.5
    if abs(y) < 0.42:
        r = y * y
        x = y * (((a[3]*r + a[2])*r + a[1])*r + a[0]) / ((((b[3]*r + b[2])*r + b[1])*r + b[0])*r + 1.0)
    else:
        r = p if y < 0 else 1.0 - p
        r = math.log(-math.log(r))
        x = c[0] + r*(c[1] + r*(c[2] + r*(c[3] + r*(c[4] + r*(c[5] + r*(c[6] + r*(c[7] + r*c[8])))))))
        if y < 0:
            x = -x

    return mu + sigma * x

def norm_pdf(x, mu=0.0, sigma=1.0):
    """정규분포 확률밀도함수(PDF)"""
    return (1.0 / (sigma * math.sqrt(2.0 * math.pi))) * math.exp(-0.5 * ((x - mu) / sigma) ** 2)

# --- 웹 환경 설정 ---
st.set_page_config(
    page_title="G2B 정밀 투찰가 분석 & 통계 솔루션",
    page_icon="🎯",
    layout="wide"
)

SERVICE_KEY = "8098f114d601ef88d54883338b215f1ad17615f330cb0b4f0a1abb20c474b777"

# --- 1. 나라장터 API 공고 수집 ---
def fetch_g2b_bid_info(bid_no_full: str):
    clean_no = bid_no_full.strip()
    if "-" in clean_no:
        bid_ntce_no, bid_ntce_ord = clean_no.split("-", 1)
    else:
        bid_ntce_no = clean_no
        bid_ntce_ord = "00"
        
    url = (
        "https://apis.data.go.kr/1230000/ad/BidPublicInfoService/getBidPblancListInfoServc"
        f"?serviceKey={SERVICE_KEY}&pageNo=1&numOfRows=1&inqryDiv=2"
        f"&bidNtceNo={bid_ntce_no}"
    )

    try:
        res = requests.get(url, timeout=10)
        if res.status_code != 200:
            return None
        root = ET.fromstring(res.text)
        item = root.find(".//item")
        if item is None:
            return None

        title = item.findtext("bidNtceNm", "공고명 없음")
        instt_nm = item.findtext("ntceInsttNm", "발주기관 없음")
        raw_base = item.findtext("bscAmt", "0").strip()
        raw_bdgt = item.findtext("bdgtAmt", "0").strip()
        raw_presmpt = item.findtext("presmptPrce", "0").strip()
        raw_a = item.findtext("exptPrfrmrestrctAmt", "0").strip()

        base_amt = 0.0
        if raw_base and float(raw_base) > 0:
            base_amt = float(raw_base)
        elif raw_bdgt and float(raw_bdgt) > 0:
            base_amt = float(raw_bdgt)
        elif raw_presmpt and float(raw_presmpt) > 0:
            base_amt = float(raw_presmpt)

        val_a = float(raw_a) if raw_a else 0.0

        return {
            "bid_no": f"{bid_ntce_no}-{bid_ntce_ord}",
            "title": title,
            "agency": instt_nm,
            "base_amt": base_amt,
            "val_a": val_a
        }
    except Exception:
        return None

# --- 2. 발주기관 최근 개찰결과 사정율 통계 수집 ---
def fetch_agency_history_stats(agency_name: str):
    if not agency_name or agency_name == "발주기관 없음":
        return None

    url = (
        "https://apis.data.go.kr/1230000/ad/OpenBiddResultInfoService/getOpengResultListInfoServc"
        f"?serviceKey={SERVICE_KEY}&pageNo=1&numOfRows=20&inqryDiv=2"
        f"&dminsttNm={requests.utils.quote(agency_name)}"
    )

    try:
        res = requests.get(url, timeout=8)
        if res.status_code != 200:
            return None
        root = ET.fromstring(res.text)
        items = root.findall(".//item")
        
        rates = []
        for item in items:
            raw_plnd = item.findtext("plndprcRate", "").strip()
            if raw_plnd:
                try:
                    val = float(raw_plnd)
                    if 95.0 <= val <= 105.0:
                        rates.append(val)
                except ValueError:
                    continue

        if not rates:
            return None

        rates_arr = np.array(rates[:15])
        return {
            "count": len(rates_arr),
            "mean": float(np.mean(rates_arr)),
            "min": float(np.min(rates_arr)),
            "max": float(np.max(rates_arr)),
            "low_bias_pct": float(np.sum(rates_arr < 100.0) / len(rates_arr) * 100)
        }
    except Exception:
        return None

# --- 3. 끝자리 7원 상향 보정 ---
def adjust_to_7(amount: float) -> int:
    amt_ceil = math.ceil(amount)
    remainder = amt_ceil % 10
    if remainder <= 7:
        return amt_ceil + (7 - remainder)
    else:
        return amt_ceil + (17 - remainder)

# --- 4. 정규분포 곡선 및 차트 ---
def build_distribution_chart(central_rate, sigma, min_rate, max_rate, peak_rate, safe_rate, att_rate, hist_mean=None):
    x_vals = np.linspace(min_rate - 0.005, max_rate + 0.005, 400)
    y_vals = [norm_pdf(x, central_rate, sigma) for x in x_vals]
    
    fig = go.Figure()

    fig.add_trace(go.Scatter(
        x=x_vals * 100, y=y_vals,
        mode='lines',
        name='사정율 확률분포 밀도',
        line=dict(color='#3B82F6', width=3),
        fill='tozeroy',
        fillcolor='rgba(59, 130, 246, 0.1)'
    ))

    points = [
        ("🏆 피크코어", peak_rate, "#10B981", 12),
        ("⚡ 공격형", att_rate, "#F59E0B", 10),
        ("🛡️ 안전형", safe_rate, "#6366F1", 10),
    ]

    for label, r_val, color, size in points:
        y_pt = norm_pdf(r_val, central_rate, sigma)
        fig.add_trace(go.Scatter(
            x=[r_val * 100], y=[y_pt],
            mode='markers+text',
            name=label,
            text=[f"{label}<br>{r_val*100:.3f}%"],
            textposition="top center",
            marker=dict(color=color, size=size, symbol='diamond')
        ))

    if hist_mean:
        fig.add_vline(
            x=hist_mean, 
            line_dash="dot", 
            line_color="#8B5CF6", 
            annotation_text=f"발주처과거평균 ({hist_mean:.3f}%)",
            annotation_position="bottom right"
        )

    fig.add_vline(x=min_rate * 100, line_dash="dash", line_color="#EF4444", annotation_text="최저하한선")
    fig.add_vline(x=max_rate * 100, line_dash="dash", line_color="#EF4444", annotation_text="최고상한선")

    fig.update_layout(
        title="📈 사정율 정규분포 곡선 및 추천 투찰 위치 시각화",
        xaxis_title="조정 사정율 (%)",
        yaxis_title="밀도(발생 확률)",
        template="plotly_white",
        height=380,
        margin=dict(l=20, r=20, t=50, b=20),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
    )
    return fig

# --- 5. 엑셀 생성 함수 ---
def create_excel_bytes(bid_info, df, summary_info):
    wb = Workbook()
    ws = wb.active
    ws.title = "입찰분석결과"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="맑은 고딕", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="맑은 고딕", size=10, bold=True)
    font_regular = Font(name="맑은 고딕", size=10)

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_meta_lbl = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"나라장터 입찰분석: {bid_info['title']} (끝자리 7원 특화)"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    meta = [
        ("공고번호 / 발주처", f"{bid_info['bid_no']} / {bid_info['agency']}"),
        ("기초금액", summary_info["base_amt"]),
        ("사후정산 A값", summary_info["val_a"]),
        ("A차감 순수모수(기초-A)", summary_info["net_base"]),
        ("참여업체수(N) / 규격", f"{summary_info['firm_count']}개사 / {summary_info['range_str']}"),
        ("적용 낙찰하한율", summary_info["lower_limit"]),
        ("목표 중앙사정율(μ)", summary_info["central_rate"])
    ]

    for idx, (label, val) in enumerate(meta, start=2):
        ws.cell(row=idx, column=1, value=label).fill = fill_meta_lbl
        ws.cell(row=idx, column=1).font = font_bold
        c_val = ws.cell(row=idx, column=2, value=val)
        c_val.font = font_bold

        if label in ["기초금액", "사후정산 A값", "A차감 순수모수(기초-A)"]:
            c_val.number_format = "#,##0"
        elif label == "적용 낙찰하한율":
            c_val.number_format = "0.000%"
        elif label == "목표 중앙사정율(μ)":
            c_val.number_format = "0.0000%"
            c_val.font = Font(name="맑은 고딕", size=10, bold=True, color="B45309")
        ws.row_dimensions[idx].height = 20

    headers = ["순위(i)", "누적분위", "조정 사정율", "A차감 기초모수", "추정예가(하한적용)", "실투찰금액(추정예가+A)", "실투찰률", "편차(vs μ)", "전략구간"]
    for col_idx, text in enumerate(headers, start=1):
        c = ws.cell(row=10, column=col_idx, value=text)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[10].height = 25

    row_cursor = 11
    for _, r in df.iterrows():
        ws.cell(row=row_cursor, column=1, value=r["순위"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_cursor, column=2, value=r["누적분위"] / 100).number_format = "0.00%"
        ws.cell(row=row_cursor, column=3, value=r["조정 사정율"] / 100).number_format = "0.0000%"
        ws.cell(row=row_cursor, column=4, value=r["A차감 기초모수"]).number_format = "#,##0"
        ws.cell(row=row_cursor, column=5, value=r["추정예가(하한적용)"]).number_format = "#,##0"
        
        c_bid = ws.cell(row=row_cursor, column=6, value=r["실투찰금액(끝자리 7)"])
        c_bid.number_format = "#,##0"
        c_bid.font = Font(name="맑은 고딕", size=10, bold=True)
        
        ws.cell(row=row_cursor, column=7, value=r["실투찰률"] / 100).number_format = "0.000%"
        ws.cell(row=row_cursor, column=8, value=r["편차(vs μ)"] / 100).number_format = "+0.0000%;-0.0000%;0.0000%"
        
        c_zone = ws.cell(row=row_cursor, column=9, value=r["전략구간"])
        c_zone.alignment = Alignment(horizontal="center")

        if "★" in str(r["전략구간"]):
            fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            for c_idx in range(1, 10):
                ws.cell(row=row_cursor, column=c_idx).fill = fill_highlight

        for c_idx in range(1, 10):
            ws.cell(row=row_cursor, column=c_idx).border = thin_border
            if not ws.cell(row=row_cursor, column=c_idx).font.bold:
                ws.cell(row=row_cursor, column=c_idx).font = font_regular

        ws.row_dimensions[row_cursor].height = 20
        row_cursor += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            max_len = max(max_len, len(val_str.encode('euc-kr', 'ignore')))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 25

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

# --- 6. UI 메인 레이아웃 ---
st.title("🎯 나라장터 입찰 투찰가 정밀 분석 & 통계 솔루션")
st.caption("A값 선차감 역산 공식 | 발주처 과거 개찰 통계 연동 | 끝자리 [7원] 상향 특화")
st.markdown("---")

with st.sidebar:
    st.header("⚙️ 분석 조건 설정")
    bid_no_input = st.text_input("공고번호 입력", placeholder="예: 20240101234-00")
    
    range_type = st.radio(
        "복수예비가격 범위 선택",
        options=["±2.0% (지자체/조달청)", "±3.0% (한전/LH/공기업)"],
        index=0
    )
    is_3pct = "3.0%" in range_type

    firm_count = st.slider("예상 참여업체수 (N)", min_value=10, max_value=150, value=50, step=5)
    
    default_mu = 100.05 if is_3pct else 99.90
    central_rate_pct = st.number_input(
        "목표 중앙사정율 (%)",
        min_value=96.0,
        max_value=104.0,
        value=default_mu,
        step=0.01,
        format="%.4f"
    )

    btn_search = st.button("🚀 공고 조회 및 투찰가 계산", type="primary", use_container_width=True)

if btn_search:
    if not bid_no_input:
        st.warning("공고번호를 입력해 주세요.")
    else:
        with st.spinner("나라장터 API 실시간 데이터 조회 중..."):
            bid_info = fetch_g2b_bid_info(bid_no_input)

        if not bid_info:
            st.error("❌ 해당 공고를 찾을 수 없습니다. 공고번호를 다시 확인해 주세요.")
        else:
            base_amt = bid_info["base_amt"]
            val_a = bid_info["val_a"]

            if base_amt == 0:
                st.warning("⚠️ 기초금액이 아직 공개되지 않은 공고입니다.")
            else:
                if base_amt >= 500_000_000:
                    lower_limit_rate = 0.85995
                    scale_tier = 1.1
                elif base_amt >= 200_000_000:
                    lower_limit_rate = 0.86995
                    scale_tier = 1.0
                else:
                    lower_limit_rate = 0.87945
                    scale_tier = 0.9

                if is_3pct:
                    min_rate, max_rate, scale_factor = 0.97, 1.03, 1.25
                    range_str = "±3.0% (97~103%)"
                    delta = 0.004
                else:
                    min_rate, max_rate, scale_factor = 0.98, 1.02, 1.00
                    range_str = "±2.0% (98~102%)"
                    delta = 0.002

                central_rate = central_rate_pct / 100.0
                sigma = ((0.7 + (4.0 / firm_count)) * scale_tier * scale_factor) / 100.0
                p_min = norm_cdf(min_rate, mu=central_rate, sigma=sigma)
                p_max = norm_cdf(max_rate, mu=central_rate, sigma=sigma)
                net_base = base_amt - val_a

                col1, col2, col3, col4 = st.columns(4)
                col1.metric("기초금액", f"{base_amt:,.0f} 원")
                col2.metric("사후정산 A값", f"{val_a:,.0f} 원")
                col3.metric("A차감 순수모수", f"{net_base:,.0f} 원")
                col4.metric("적용 낙찰하한율", f"{lower_limit_rate * 100:.3f} %")

                st.info(f"**공고명:** {bid_info['title']}  |  **발주기관:** {bid_info['agency']}")

                with st.spinner(f"'{bid_info['agency']}' 최근 개찰결과 사정율 통계 분석 중..."):
                    agency_stats = fetch_agency_history_stats(bid_info['agency'])

                hist_mean_val = None
                if agency_stats:
                    hist_mean_val = agency_stats['mean']
                    st.markdown("#### 🏛️ 발주기관 과거 낙찰 사정율 통계 분석")
                    sc1, sc2, sc3, sc4 = st.columns(4)
                    sc1.metric("분석 표본 건수", f"최근 {agency_stats['count']}건")
                    sc2.metric("발주처 평균 사정율", f"{agency_stats['mean']:.4f} %")
                    sc3.metric("사정율 변동폭", f"{agency_stats['min']:.3f}% ~ {agency_stats['max']:.3f}%")
                    sc4.metric("100% 미만(저가) 낙찰률", f"{agency_stats['low_bias_pct']:.1f} %")
                    
                    if abs(agency_stats['mean'] - central_rate_pct) > 0.1:
                        st.caption(f"💡 팁: 현재 입력하신 목표 사정율({central_rate_pct:.3f}%)과 발주처 실제 평균({agency_stats['mean']:.3f}%) 간 차이가 있습니다. 왼쪽 사이드바에서 목표 사정율을 조정해 보세요.")
                else:
                    st.caption("ℹ️ 발주기관의 최근 용역 개찰 이력이 부족하여 기본 표준 모델로 계산합니다.")

                rows = []
                for i in range(1, firm_count + 1):
                    p_val_raw = (i - 0.5) / firm_count
                    p_adj = p_min + p_val_raw * (p_max - p_min)
                    target_rate = norm_ppf(p_adj, mu=central_rate, sigma=sigma)
                    target_rate = max(min_rate, min(max_rate, target_rate))

                    est_price_net = round(net_base * target_rate * lower_limit_rate)
                    final_bid = adjust_to_7(est_price_net + val_a)
                    deviation = target_rate - central_rate

                    if target_rate < central_rate - delta:
                        zone_tag = "저가구간"
                    elif target_rate < central_rate:
                        zone_tag = "★ 중앙- (피크코어)"
                    elif target_rate < central_rate + delta:
                        zone_tag = "★ 중앙+ (피크코어)"
                    else:
                        zone_tag = "고가구간"

                    rows.append({
                        "순위": i,
                        "누적분위": round(p_val_raw * 100, 2),
                        "조정 사정율": round(target_rate * 100, 4),
                        "A차감 기초모수": int(net_base),
                        "추정예가(하한적용)": int(est_price_net),
                        "실투찰금액(끝자리 7)": int(final_bid),
                        "실투찰률": round(final_bid / base_amt * 100, 4),
                        "편차(vs μ)": round(deviation * 100, 4),
                        "전략구간": zone_tag
                    })

                df = pd.DataFrame(rows)

                mid_idx = firm_count // 2
                peak_pick = rows[mid_idx]
                att_pick = rows[max(0, mid_idx - int(firm_count * 0.15))]
                safe_pick = rows[min(firm_count - 1, mid_idx + int(firm_count * 0.15))]

                st.subheader("💡 3대 전략 추천 투찰가")
                rc1, rc2, rc3 = st.columns(3)
                
                with rc1:
                    st.success("🏆 **피크 코어 추천가 (1순위 타깃)**")
                    st.write(f"**금액:** `{peak_pick['실투찰금액(끝자리 7)']:,.0f} 원`")
                    st.caption(f"사정율: {peak_pick['조정 사정율']:.4f}% | {peak_pick['순위']}위")

                with rc2:
                    st.warning("⚡ **공격적 추천가 (마진 극대화)**")
                    st.write(f"**금액:** `{att_pick['실투찰금액(끝자리 7)']:,.0f} 원`")
                    st.caption(f"사정율: {att_pick['조정 사정율']:.4f}% | {att_pick['순위']}위")

                with rc3:
                    st.info("🛡️️ **안정형 추천가 (낙찰 확률 방어)**")
                    st.write(f"**금액:** `{safe_pick['실투찰금액(끝자리 7)']:,.0f} 원`")
                    st.caption(f"사정율: {safe_pick['조정 사정율']:.4f}% | {safe_pick['순위']}위")

                chart_fig = build_distribution_chart(
                    central_rate=central_rate,
                    sigma=sigma,
                    min_rate=min_rate,
                    max_rate=max_rate,
                    peak_rate=peak_pick['조정 사정율'] / 100.0,
                    safe_rate=safe_pick['조정 사정율'] / 100.0,
                    att_rate=att_pick['조정 사정율'] / 100.0,
                    hist_mean=hist_mean_val
                )
                st.plotly_chart(chart_fig, use_container_width=True)

                st.markdown("---")

                st.subheader(f"📊 전체 {firm_count}개 순위별 투찰 시뮬레이션")
                
                display_df = df.copy()
                display_df["누적분위"] = display_df["누적분위"].map("{:.2f}%".format)
                display_df["조정 사정율"] = display_df["조정 사정율"].map("{:.4f}%".format)
                display_df["A차감 기초모수"] = display_df["A차감 기초모수"].map("{:,}원".format)
                display_df["추정예가(하한적용)"] = display_df["추정예가(하한적용)"].map("{:,}원".format)
                display_df["실투찰금액(끝자리 7)"] = display_df["실투찰금액(끝자리 7)"].map("{:,}원".format)
                display_df["실투찰률"] = display_df["실투찰률"].map("{:.3f}%".format)
                display_df["편차(vs μ)"] = display_df["편차(vs μ)"].map("{:+.4f}%".format)

                st.dataframe(display_df, use_container_width=True, height=350)

                summary_info = {
                    "base_amt": base_amt,
                    "val_a": val_a,
                    "net_base": net_base,
                    "firm_count": firm_count,
                    "range_str": range_str,
                    "lower_limit": lower_limit_rate,
                    "central_rate": central_rate
                }
                excel_bytes = create_excel_bytes(bid_info, df, summary_info)
                
                st.download_button(
                    label="📥 정밀 분석 엑셀 보고서 다운로드 (.xlsx)",
                    data=excel_bytes,
                    file_name=f"입찰분석_{bid_info['bid_no']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )
else:
    st.info("👈 왼쪽 사이드바에 공고번호를 입력하고 **[공고 조회 및 투찰가 계산]** 버튼을 눌러보세요.")
